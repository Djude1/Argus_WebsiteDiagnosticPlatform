from pathlib import Path
from collections import Counter
from zipfile import ZipFile
import json
import hashlib
import re
import ast
from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from lxml import etree

HERE = Path(__file__).resolve().parents[2] / '本機資料' / 'work' / 'v10'
ROOT = HERE.parents[1]
final = ROOT / 'Argus_系統手冊_v10_複審精簡版.docx'
v4 = ROOT / 'Argus_系統手冊_v4.docx'
v9 = ROOT / 'Argus_系統手冊_v9_新版整合修訂版.docx'
d = Document(final)
s = Document(v9)
r = Document(v4)
errors = []

def clean(text):
    return re.sub(r'\s+', '', text.replace('\u200b', ''))

def text(e):
    return ''.join(e.xpath('.//w:t/text()'))

def table_records(doc, pattern):
    out = []
    cap = ''
    for e in doc.element.body:
        if e.tag == qn('w:p'):
            cap = text(e).lstrip('▼')
        elif e.tag == qn('w:tbl') and re.match(pattern, cap):
            t = Table(e, doc._body)
            out.append((cap, [[clean(c.text) for c in row.cells] for row in t.rows]))
    return out

src = table_records(s, r'^表8-2-')
out = table_records(d, r'^表8-2-')
if [x[1] for x in src] != [x[1] for x in out]:
    errors.append('資料字典欄位內容與 v9 不一致')
fields = sum(len(rows) - 1 for _, rows in out)
if len(out) != 40 or fields != 450:
    errors.append('模型或欄位数量不符')

# 技術內容逐段核對，允許本次明確的圖題重編、附錄移動及來源修正。
all_text = clean(text(d.element.body))
allowed = {855, 858, 860, 861, 862, 863, 864, 866, 868, 870}
missing = []
for i, e in enumerate(s.element.body):
    if not (263 <= i < 279 or 306 <= i < 872) or e.tag != qn('w:p') or i in allowed:
        continue
    t = text(e)
    if not t.strip() or re.match(r'^(?:圖|表)\d+-', t):
        continue
    if 425 <= i <= 609 and (t.startswith('資料表 ') or e.xpath('./w:pPr/w:pStyle[@w:val="Heading3"]')):
        continue
    if clean(t) not in all_text:
        missing.append({'index': i, 'text': t[:160]})
if missing:
    errors.append('有未核准的技術段落差異')

names = d.element.xpath('.//w:bookmarkStart/@w:name')
ids = d.element.xpath('.//w:bookmarkStart/@w:id')
links = d.element.xpath('.//w:hyperlink/@w:anchor')
if len(names) != len(set(names)) or len(ids) != len(set(ids)):
    errors.append('重複書籤')
if set(links) - set(names):
    errors.append('失效目錄連結')
docpr = d.element.xpath('.//wp:docPr/@id')
if len(docpr) != len(set(docpr)):
    errors.append('圖片 ID 重複')
for e in d.element.iter():
    for attr in (qn('r:embed'), qn('r:id'), qn('r:link')):
        rid = e.get(attr)
        if rid and rid not in d.part.rels:
            errors.append('缺少關係 ' + rid)
manifest = json.loads((HERE / 'manifest.json').read_text(encoding='utf-8'))
mapping = json.loads((HERE / 'page_map.json').read_text(encoding='utf-8-sig'))
if set(x['bookmark'] for x in manifest) - set(mapping):
    errors.append('頁碼對映缺項')

ai = [t for t in d.tables if any('OpenAI Codex' in c.text.replace('\u200b', '') for row in t.rows for c in row.cells)]
if len(ai) != 1 or len(ai[0].rows) != 8:
    errors.append('AI 使用說明表不完整')
if '附錄一人工智慧使用說明表' not in all_text:
    errors.append('缺少指定附錄一')
if '錯誤!尚未定義書籤' in all_text or '[序號' in all_text:
    errors.append('殘留舊版錯誤或序號註記')
if hashlib.sha256(v4.read_bytes()).hexdigest() != '739c11b8d9869bf1955d8590561ede09241fd2206e4d83e0b2ce56897f07b0c6':
    errors.append('v4 原稿被更改')
if hashlib.sha256(v9.read_bytes()).hexdigest() != '7d22674fe9c099d9db30ee01d9853ea80463371f210f25c465c8b67027719e49':
    errors.append('v9 原稿被更改')

# 保留主題、編號及原始媒體的內容雜湊。
with ZipFile(v4) as z, ZipFile(final) as f:
    preserved = []
    for name in z.namelist():
        if name.startswith(('word/theme/', 'word/media/')) or name == 'word/numbering.xml':
            if name not in f.namelist() or z.read(name) != f.read(name):
                errors.append('參考保存部分被更動 ' + name)
            preserved.append(name)

for file in Path(__file__).resolve().parent.glob('*.py'):
    ast.parse(file.read_text(encoding='utf-8'))
report = {'errors': errors, 'missing_technical_paragraphs': missing, 'models': len(out), 'fields': fields, 'figures': sum(x['kind'] == 'figure' for x in manifest), 'tables': len(d.tables), 'catalog_targets': len(manifest), 'ai_rows': len(ai[0].rows) - 1 if ai else 0, 'preserve_only_parts': len(preserved), 'sha256': hashlib.sha256(final.read_bytes()).hexdigest()}
(HERE / 'final_audit.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps(report, ensure_ascii=False, indent=2))
raise SystemExit(bool(errors))
