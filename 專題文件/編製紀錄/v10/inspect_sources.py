from pathlib import Path
import hashlib
import json
from collections import Counter
from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2] / '本機資料'
HERE = Path(__file__).resolve().parents[2] / '本機資料' / 'work' / 'v10'

def inspect(path, label):
    doc = Document(path)
    records = []
    for i, element in enumerate(doc.element.body):
        tag = element.tag.rsplit('}', 1)[-1]
        if tag == 'p':
            text = ''.join(element.xpath('.//w:t/text()'))
        elif tag == 'tbl':
            text = '\n'.join(' | '.join(''.join(c.xpath('.//w:t/text()')) for c in row.xpath('./w:tc')) for row in element.xpath('./w:tr'))
        else:
            text = ''
        styles = element.xpath('./w:pPr/w:pStyle/@w:val') if tag in ('p', 'tbl') else []
        records.append({'i': i, 'tag': tag, 'style': styles[0] if styles else '', 'text': text, 'images': element.xpath('.//a:blip/@r:embed') if tag in ('p', 'tbl') else []})
    (HERE / f'{label}_blocks.json').write_text(json.dumps(records, ensure_ascii=False, indent=2), encoding='utf-8')
    (HERE / f'{label}_complete.txt').write_text('\n\n'.join(f'[{r["i"]} {r["tag"]} {r["style"]}]\n{r["text"]}' for r in records), encoding='utf-8')
    sections = []
    for s in doc.sections:
        sections.append({'width_cm': s.page_width.cm, 'height_cm': s.page_height.cm, 'margins_cm': [s.top_margin.cm, s.right_margin.cm, s.bottom_margin.cm, s.left_margin.cm], 'header_cm': s.header_distance.cm, 'footer_cm': s.footer_distance.cm, 'start': str(s.start_type)})
    styles = {}
    for s in doc.styles:
        if s.type == 1:
            styles[s.style_id] = {'name': s.name, 'xml': s.element.xml}
    evidence = {'sha256': hashlib.sha256(path.read_bytes()).hexdigest(), 'sections': sections, 'styles': styles, 'fonts': Counter(doc.element.xpath('.//w:rFonts/@w:eastAsia')), 'colors': Counter(doc.element.xpath('.//w:color/@w:val')), 'fills': Counter(doc.element.xpath('.//w:shd/@w:fill')), 'tables': len(doc.tables), 'images': len(doc.inline_shapes)}
    (HERE / f'{label}_evidence.json').write_text(json.dumps(evidence, ensure_ascii=False, indent=2), encoding='utf-8')
    print(label, 'blocks', len(records), 'tables', len(doc.tables), 'inline_images', len(doc.inline_shapes), 'sections', sections)
    print('headings', [(r['i'], r['style'], r['text']) for r in records if r['style'].lower().startswith('heading') or r['style'] in ('1', '2', '3')])
    print('fonts', evidence['fonts'], 'colors', evidence['colors'], 'fills', evidence['fills'])

inspect(ROOT / 'Argus_系統手冊_v4.docx', 'v4')
inspect(ROOT / 'Argus_系統手冊_v9_新版整合修訂版.docx', 'v9')
