"""OWASP Top 10 2025 對照與被動 API 檢查（2026-10-11，回應評審 2-3、3-7）。"""

from __future__ import annotations

import asyncio
from types import SimpleNamespace

from django.contrib.auth import get_user_model
from django.test import SimpleTestCase, TestCase
from rest_framework.test import APIClient

from apps.scans import crawler, tasks
from apps.scans.coverage import COMPLETED, SKIPPED
from apps.scans.models import Finding, ScanJob
from apps.scans.reports import _owasp_cwe_text
from apps.scans.scan_plan import build_scan_execution_plan
from apps.scans.security import owasp_mapper
from apps.scans.security.api_scanner import MAX_SURFACE_ENDPOINTS, analyze_api
from apps.scans.security.finding_kind import security_kind
from apps.scans.serializers import FindingSerializer

User = get_user_model()

DJANGO_TRACE = (
    "Traceback (most recent call last):\n  File \"/srv/app/views.py\", line 42, in get\n"
    "    raise ValueError('owner@example.com')\nValueError"
)


def _obs(path="/api/items", method="GET", status=200, content_type="application/json", **kw):
    headers = {"content-type": content_type}
    headers.update(kw.pop("headers", {}))
    return {"method": method, "path": path, "status": status, "headers": headers, **kw}


class Owasp2025MappingTests(SimpleTestCase):
    def test_2021_codes_map_to_2025(self):
        self.assertEqual(owasp_mapper.owasp_2025("A05"), "A02")  # 安全設定錯誤
        self.assertEqual(owasp_mapper.owasp_2025("A06"), "A03")  # 供應鏈
        self.assertEqual(owasp_mapper.owasp_2025("A03"), "A05")  # 注入
        self.assertEqual(owasp_mapper.owasp_2025("A10"), "A01")  # SSRF 併入存取控制
        self.assertEqual(owasp_mapper.owasp_2025(""), "")

    def test_cwe_209_goes_to_exceptional_conditions(self):
        self.assertEqual(owasp_mapper.owasp_2025("A05", "CWE-209"), "A10")

    def test_every_2021_code_has_a_2025_label(self):
        for code in owasp_mapper.OWASP_2021_TO_2025.values():
            self.assertIn(code, owasp_mapper.OWASP_2025_LABELS)

    def test_api_rules_are_mapped(self):
        tagged = owasp_mapper.tag({"category": "security", "rule_id": "api-error-stack-trace"})
        self.assertEqual((tagged["owasp_category"], tagged["cwe_id"]), ("A05", "CWE-209"))
        kind = lambda rule: security_kind(category="security", rule_id=rule)  # noqa: E731
        self.assertEqual(kind("api-error-stack-trace"), "exposure")
        self.assertEqual(kind("api-cors-wildcard"), "config")


class Owasp2025DisplayTests(TestCase):
    def setUp(self):
        user = User.objects.create_user(username="owasp25", password="safe-test-password")
        scan = ScanJob.objects.create(
            user=user, original_url="https://example.com/", normalized_url="https://example.com/",
            origin="https://example.com",
        )
        self.finding = Finding.objects.create(
            scan_job=scan, category="security", severity="medium", title="缺少 HSTS",
            description="d", remediation="r", rule_id="x", owasp_category="A05", cwe_id="CWE-319",
        )

    def test_serializer_and_report_show_2025_with_2021(self):
        data = FindingSerializer(self.finding).data
        self.assertEqual(data["owasp_2025"], "A02")
        self.assertIn("安全設定錯誤", data["owasp_2025_label"])
        self.assertEqual(_owasp_cwe_text(self.finding), "A02:2025（2021 A05） / CWE-319")


class AnalyzeApiTests(SimpleTestCase):
    def test_stack_trace_in_error_response_is_reported_and_redacted(self):
        findings, surface = analyze_api(
            [_obs(status=500, error_body=DJANGO_TRACE), _obs(path="/api/ok")]
        )
        self.assertEqual([f["rule_id"] for f in findings], ["api-error-stack-trace"])
        self.assertEqual(findings[0]["severity"], "medium")
        self.assertIn("Python", findings[0]["evidence"])
        self.assertNotIn("owner@example.com", findings[0]["evidence"])
        self.assertEqual(surface["total"], 2)
        self.assertEqual(surface["with_issues"], 1)
        self.assertEqual(surface["endpoints"][0]["rules"], ["api-error-stack-trace"])

    def test_plain_error_message_is_not_a_leak(self):
        findings, _ = analyze_api([_obs(status=404, error_body='{"detail": "Not found."}')])
        self.assertEqual(findings, [])

    def test_wildcard_cors_only_for_json(self):
        star = {"access-control-allow-origin": "*"}
        findings, _ = analyze_api([
            _obs(headers=star),
            _obs(path="/static/a.css", content_type="text/css", headers=star),
        ])
        self.assertEqual([f["rule_id"] for f in findings], ["api-cors-wildcard"])
        self.assertIn("GET /api/items", findings[0]["evidence"])

    def test_surface_is_capped(self):
        _, surface = analyze_api([_obs(path=f"/api/{i}") for i in range(MAX_SURFACE_ENDPOINTS + 5)])
        self.assertEqual(surface["total"], MAX_SURFACE_ENDPOINTS + 5)
        self.assertEqual(len(surface["endpoints"]), MAX_SURFACE_ENDPOINTS)


class CrawlerObservationTests(SimpleTestCase):
    def _resp(self, url, status=200, headers=None, method="GET", body=""):
        async def text():
            return body

        return SimpleNamespace(
            url=url, status=status, headers=headers or {"content-type": "application/json"},
            request=SimpleNamespace(method=method), text=text,
        )

    def test_records_path_without_query_and_reads_error_body(self):
        sink: dict = {}
        tasks_list: list = []

        async def run():
            crawler._observe_api_response(
                self._resp("https://example.com/api/a?token=secret", status=500, body=DJANGO_TRACE),
                sink, tasks_list,
            )
            crawler._observe_api_response(
                self._resp("https://example.com/api/a?token=other"), sink, tasks_list
            )
            await crawler._finish_api_body_tasks(tasks_list)

        asyncio.run(run())
        self.assertEqual(list(sink), ["GET /api/a"])
        obs = sink["GET /api/a"]
        self.assertEqual(obs["path"], "/api/a")
        self.assertNotIn("secret", str(obs["headers"]))
        self.assertIn("Traceback", obs["error_body"])

    def test_cloudflare_endpoints_are_not_site_apis(self):
        sink: dict = {}
        crawler._observe_api_response(self._resp("https://example.com/cdn-cgi/rum?x"), sink, [])
        self.assertEqual(sink, {})

    def test_success_response_body_is_not_read(self):
        sink: dict = {}
        tasks_list: list = []
        crawler._observe_api_response(self._resp("https://example.com/api/b"), sink, tasks_list)
        self.assertEqual(tasks_list, [])
        self.assertNotIn("error_body", sink["GET /api/b"])


class PassiveApiStageTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username="api-stage", password="safe-test-password")
        self.scan = ScanJob.objects.create(
            user=self.user, original_url="https://example.com/",
            normalized_url="https://example.com/", origin="https://example.com",
            status=ScanJob.Status.COMPLETED, max_pages=1, max_depth=1, categories=["security"],
        )
        self.ctx = tasks.ScanRunContext(
            scan_job=self.scan, execution_plan=build_scan_execution_plan(self.scan), steps=[],
            crawl_phase_started="",
        )

    def test_no_observations_is_skipped(self):
        tasks._passive_api_check(self.ctx)
        self.assertEqual(self.ctx.coverage.status_of("api_passive"), SKIPPED)

    def test_security_page_reports_check_status_without_surface(self):
        self.scan.coverage = {"checks": {"api_passive": {"status": "skipped", "reason": ""}}}
        self.scan.save(update_fields=["coverage"])
        client = APIClient()
        client.force_authenticate(user=self.user)
        data = client.get(f"/api/projects/{self.scan.project_id}/security/").json()
        self.assertIsNone(data["api_surface"])
        self.assertEqual(data["api_check"]["status"], "skipped")

    def test_findings_recorded_and_surface_saved_then_shown_on_security_page(self):
        self.ctx.site_signals["api_observations"] = [_obs(status=500, error_body=DJANGO_TRACE)]
        tasks._passive_api_check(self.ctx)
        self.assertEqual(self.ctx.coverage.status_of("api_passive"), COMPLETED)
        finding = Finding.objects.get(scan_job=self.scan, rule_id="api-error-stack-trace")
        self.assertEqual(finding.cwe_id, "CWE-209")
        self.assertEqual(finding.owasp_2025, "A10")
        self.scan.refresh_from_db()
        self.assertEqual(self.scan.warning_summary["api_surface"]["total"], 1)

        client = APIClient()
        client.force_authenticate(user=self.user)
        data = client.get(f"/api/projects/{self.scan.project_id}/security/").json()
        self.assertEqual(data["api_surface"]["endpoints"][0]["path"], "/api/items")

    def test_static_js_paths_surface_without_runtime_calls(self):
        # 公開頁不呼叫後端 API（無觀察），但前端 JS 靜態抽出端點 → 仍要列出、不判 skipped
        self.ctx.site_signals["js_api_paths"] = ["/gateway", "/api/v1/users"]
        tasks._passive_api_check(self.ctx)
        self.assertEqual(self.ctx.coverage.status_of("api_passive"), COMPLETED)
        self.scan.refresh_from_db()
        surface = self.scan.warning_summary["api_surface"]
        self.assertEqual(surface["total"], 0)
        self.assertEqual(surface["static_discovered"], 2)
        rows = surface["endpoints"]
        self.assertTrue(all(r["source"] == "js_static" and r["method"] == "靜態" for r in rows))
        self.assertEqual({r["path"] for r in rows}, {"/gateway", "/api/v1/users"})

    def test_static_paths_exclude_crawled_pages(self):
        # 爬蟲當網頁造訪過的路徑不是 API（2026-10-17 實掃：/service/1 被列進 API 清單）
        self.ctx.crawled_pages = [{"url": "https://example.com/service/1",
                                   "final_url": "https://example.com/service/1"}]
        self.ctx.site_signals["js_api_paths"] = ["/gateway", "/service/1"]
        tasks._passive_api_check(self.ctx)
        self.scan.refresh_from_db()
        paths = [r["path"] for r in self.scan.warning_summary["api_surface"]["endpoints"]]
        self.assertEqual(paths, ["/gateway"])
        self.assertEqual(tasks._observed_api_paths(self.ctx), ["/gateway"])

    def test_static_paths_do_not_duplicate_observed(self):
        self.ctx.site_signals["api_observations"] = [_obs(path="/api/items")]
        self.ctx.site_signals["js_api_paths"] = ["/api/items", "/gateway"]
        tasks._passive_api_check(self.ctx)
        self.scan.refresh_from_db()
        surface = self.scan.warning_summary["api_surface"]
        # /api/items 已被觀察到，不再以靜態重複列入
        static_rows = [r for r in surface["endpoints"] if r.get("source") == "js_static"]
        self.assertEqual([r["path"] for r in static_rows], ["/gateway"])


class ApiSurfaceReportTests(TestCase):
    """PDF 報告附錄「API 端點清單」（2026-10-12）。"""

    def setUp(self):
        user = User.objects.create_user(username="api-report", password="safe-test-password")
        self.scan = ScanJob.objects.create(
            user=user, original_url="https://example.com/", normalized_url="https://example.com/",
            origin="https://example.com", status=ScanJob.Status.COMPLETED, categories=["security"],
            warning_summary={"api_surface": {"total": 2, "with_issues": 1, "endpoints": [
                {"method": "GET", "path": "/api/a", "status": 500, "content_type": "text/plain",
                 "rules": ["api-error-stack-trace"]},
                {"method": "POST", "path": "/api/b", "status": 200, "content_type": "",
                 "rules": []},
            ]}},
        )

    def test_payload_has_rows_with_labels(self):
        from apps.scans.reports import _report_api_surface

        api = _report_api_surface(self.scan)
        self.assertEqual(api["total"], 2)
        self.assertEqual(api["rows"][0]["issues"], "錯誤回應洩漏堆疊")
        self.assertEqual(api["rows"][1]["content_type"], "—")

    def test_static_discovered_row_labelled_in_report(self):
        from apps.scans.reports import _report_api_surface

        self.scan.warning_summary = {"api_surface": {
            "total": 0, "with_issues": 0, "static_discovered": 1,
            "endpoints": [{"method": "靜態", "path": "/gateway", "status": None,
                           "content_type": "", "rules": [], "source": "js_static"}],
        }}
        api = _report_api_surface(self.scan)
        self.assertEqual(api["static_discovered"], 1)
        self.assertEqual(api["rows"][0]["issues"], "前端 JS 靜態發現，執行期未呼叫")
        self.assertEqual(api["rows"][0]["status"], "—")

    def test_docx_contains_section(self):
        import tempfile
        from pathlib import Path

        from docx import Document

        from apps.scans.reports import render_report_docx

        with tempfile.TemporaryDirectory() as workdir:
            path = Path(workdir) / "report.docx"
            render_report_docx(self.scan, path)
            doc = Document(str(path))
        text = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("API 端點清單", text)
        cells = {c.text for t in doc.tables for r in t.rows for c in r.cells}
        self.assertIn("/api/a", cells)

    def test_not_listed_when_security_not_checked(self):
        from apps.scans.reports import _report_api_surface

        self.scan.categories = ["seo"]
        self.assertIsNone(_report_api_surface(self.scan))
