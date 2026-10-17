"""報告第 1 章「網站總覽」裡的效能、連結、搜尋詞與拓撲（2026-10-11，RENDERER_VERSION 23／25）。

網頁上有、報告原本沒有的資訊：PageSpeed 效能分數、連結檢查、Search Console 搜尋詞、網站拓撲。
"""

from __future__ import annotations

import json
import tempfile
from pathlib import Path
from unittest import mock

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.utils import timezone
from docx import Document
from jsonschema import Draft7Validator

from apps.scans.models import Page, ScanJob, SearchConsoleConnection
from apps.scans.report_site_metrics import site_metrics
from apps.scans.reports import build_report_payload, render_report_docx

User = get_user_model()
SCHEMA = json.loads(
    (Path(__file__).resolve().parent / "report_render" / "schema.json").read_text("utf-8")
)
PERFORMANCE = {"lab": {
    "final_url": "https://example.com/",
    "scores": {"performance": 62, "accessibility": 91, "best-practices": 100, "seo": 85},
    "metrics": {"largest-contentful-paint": {"label": "LCP（最大內容繪製）", "display": "4.1 s",
                                             "score": 0.3}},
    "opportunities": [{"title": "減少未使用的 JavaScript", "display": "可節省 1.2 秒"}],
}, "field": {"scope": "none", "reason": "CrUX 資料不足"}}
GSC_DATA = {
    "property": "sc-domain:example.com", "start": "2026-09-12", "end": "2026-10-09",
    "totals": {"clicks": 80, "impressions": 3283, "ctr": 0.0244},
    "trend": [{"date": "2026-09-12", "clicks": 3, "impressions": 100},
              {"date": "2026-09-13", "clicks": 5, "impressions": 140}],
    "queries": [
        {"query": "birc", "clicks": 7, "impressions": 15, "ctr": 0.467, "position": 1.5,
         "page": "https://example.com/"},
        {"query": "聯絡 0912-345-678", "clicks": 1, "impressions": 3, "ctr": 0.33,
         "position": 4.0, "page": "https://example.com/contact"},
    ],
}


class ReportSiteMetricsTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username="metrics", password="safe-test-password")
        self.scan_job = ScanJob.objects.create(
            user=self.user, original_url="https://example.com/",
            normalized_url="https://example.com/", origin="example.com",
            status=ScanJob.Status.COMPLETED, completed_at=timezone.now(),
            overall_score=70, category_scores={"seo": 70},
        )

    def _page(self, path, depth, links=()):
        return Page.objects.create(
            scan_job=self.scan_job, url=f"https://example.com{path}",
            final_url=f"https://example.com{path}", origin="example.com",
            status_code=200, depth=depth,
            outgoing_links=[f"https://example.com{link}" for link in links],
        )

    def _text(self):
        with tempfile.TemporaryDirectory() as workdir:
            path = Path(workdir) / "report.docx"
            render_report_docx(self.scan_job, path)
            doc = Document(str(path))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts += [cell.text for cell in row.cells]
        return "\n".join(parts)

    def test_empty_scan_explains_each_missing_section(self):
        metrics = site_metrics(self.scan_job)
        self.assertIn("PageSpeed", metrics["performance"]["note"])
        self.assertIn("SEO", metrics["links"]["note"])
        self.assertIn("尚未連接 Google Search Console", metrics["search"]["note"])
        self.assertIn("沒有可以畫的網站結構", metrics["topology"]["note"])
        payload = build_report_payload(self.scan_job)
        self.assertEqual(list(Draft7Validator(SCHEMA).iter_errors(payload)), [])
        text = self._text()
        # 第 1 章「網站總覽」先呈現圖表與量測，要處理的問題從第 2 章開始
        for heading in ("1　網站總覽", "Google Lighthouse 首頁檢查（PageSpeed Insights）",
                        "2　優先處理清單", "3　這些分類為什麼重要", "5　掃描資訊與範圍", "6　附錄"):
            self.assertIn(heading, text)
        # 附錄小節連續編號（沒有名詞解釋與技術索引時從 6.1 開始），免責固定在最後
        self.assertIn("6.1　修補後如何驗證", text.split("6　附錄")[-1])
        self.assertTrue(text.split("6　附錄")[-1].rstrip().count("免責與報告產生方式"))
        self.assertLess(text.index("網站拓撲"), text.index("2　優先處理清單"))

    def test_performance_scores_and_metrics(self):
        self.scan_job.performance_report = PERFORMANCE
        self.scan_job.save()
        perf = site_metrics(self.scan_job)["performance"]
        self.assertEqual([s["score"] for s in perf["scores"]], [62, 91, 100, 85])
        # 與 Argus 分數同名的兩項標成「基本檢查」，並說明兩邊不是同一個分數
        self.assertEqual([s["label"] for s in perf["scores"]],
                         ["效能", "無障礙基本檢查", "最佳做法", "SEO 基本檢查"])
        self.assertEqual(perf["metrics"][0]["rating"], "不佳")
        self.assertEqual(perf["field_note"], "CrUX 資料不足")
        text = self._text()
        self.assertIn("Google Lighthouse 首頁檢查（PageSpeed Insights）", text)
        self.assertIn("減少未使用的 JavaScript", text)
        self.assertIn("只看首頁的基本項目；本章開頭的 Argus 分數檢查", text)

    def test_absurd_lab_lcp_gets_note(self):
        # 2026-10-17 實掃：實驗室 LCP 240.8 秒，真實使用者 1.7 秒
        import copy

        report = copy.deepcopy(PERFORMANCE)
        report["lab"]["metrics"]["largest-contentful-paint"].update(
            {"value": 240_800, "display": "240.8 s"})
        self.scan_job.performance_report = report
        self.scan_job.save()
        perf = site_metrics(self.scan_job)["performance"]
        self.assertIn("240.8 s", perf["lab_note"])
        self.assertIn(perf["lab_note"], self._text())
        # 一般慢（4.1 秒）不加說明
        self.scan_job.performance_report = PERFORMANCE
        self.scan_job.save()
        self.assertNotIn("lab_note", site_metrics(self.scan_job)["performance"])

    def test_topology_label_keeps_query_string(self):
        from apps.scans.report_site_metrics import _path

        self.assertEqual(_path("https://example.com/p/1.php?Lang=en"), "/p/1.php?Lang=en")
        self.assertEqual(_path("https://example.com"), "/")

    def test_search_console_queries_are_included_and_masked(self):
        connection = SearchConsoleConnection.objects.create(
            project=self.scan_job.project, user=self.user, property_url="sc-domain:example.com",
            refresh_token_encrypted="x",
        )
        self.assertIsNotNone(connection.pk)
        with mock.patch("apps.scans.seo.gsc.is_enabled", return_value=True), \
                mock.patch("apps.scans.seo.gsc.performance", return_value=GSC_DATA) as perf:
            search = site_metrics(self.scan_job)["search"]
            text = self._text()
        perf.assert_called()
        self.assertEqual(search["totals"], {"clicks": 80, "impressions": 3283, "ctr": 0.0244,
                                            "queries": 2})
        self.assertEqual(search["period"], "2026-09-12～2026-10-09")
        self.assertNotIn("0912-345-678", search["queries"][1]["query"])
        self.assertIn("實際搜尋詞（Google Search Console）", text)
        self.assertIn("birc", text)

    def test_search_console_failure_does_not_break_report(self):
        SearchConsoleConnection.objects.create(
            project=self.scan_job.project, user=self.user, property_url="sc-domain:example.com",
            refresh_token_encrypted="x",
        )
        with mock.patch("apps.scans.seo.gsc.is_enabled", return_value=True), \
                mock.patch("apps.scans.seo.gsc.performance", side_effect=RuntimeError("down")):
            search = site_metrics(self.scan_job)["search"]
        self.assertIn("無法取得 Search Console 資料", search["note"])

    def test_topology_nodes_and_edges(self):
        self._page("/", 0, links=["/a", "/b"])
        self._page("/a", 1, links=["/b"])
        self._page("/b", 1)
        topo = site_metrics(self.scan_job)["topology"]
        self.assertEqual([n["label"] for n in topo["nodes"]], ["/", "/a", "/b"])
        self.assertEqual(len(topo["edges"]), 3)
        self.assertIn("網站拓撲", self._text())


class ReportDetailFixesTests(TestCase):
    """2026-10-11 逐頁審查：AEO 獨立成章、免責在最後、CVE 收合、共用元件、日期格式。"""

    def setUp(self):
        self.user = User.objects.create_user(username="details", password="safe-test-password")
        self.scan_job = ScanJob.objects.create(
            user=self.user, original_url="https://example.com/",
            normalized_url="https://example.com/", origin="example.com",
            status=ScanJob.Status.COMPLETED, completed_at=timezone.now(),
            overall_score=70, category_scores={"security": 70},
        )

    def _finding(self, rule_id, title, page=None, evidence="e", description="d",
                 category="security", severity="medium"):
        from apps.scans.models import Finding

        Finding.objects.create(
            scan_job=self.scan_job, page=page, category=category, severity=severity,
            title=title, description=description, remediation="r", rule_id=rule_id,
            evidence=evidence, ai_handoff_prompt="p", priority_score=10,
        )

    def _page(self, path):
        return Page.objects.create(
            scan_job=self.scan_job, url=f"https://example.com{path}",
            final_url=f"https://example.com{path}", origin="example.com",
            status_code=200, depth=1,
        )

    def _text(self):
        return ReportSiteMetricsTests._text(self)

    def test_aeo_items_are_their_own_chapter_and_disclaimer_is_last(self):
        self.scan_job.aeo_report = {
            "status": "evaluated", "questions_total": 1,
            "counts": {"answered": 1, "insufficient": 0, "conflict": 0, "missing": 0},
            "answered_ratio": 1.0,
            "questions": [{"text": "營業時間是什麼時候？", "verdict": "answered",
                           "verdict_label": "可回答", "evidence": [
                               {"url": "https://example.com/", "location": "段落",
                                "quote": "每天 9:00–18:00"}]}],
        }
        self.scan_job.save()
        text = self._text()
        self.assertIn("5　AEO 問答檢測結果", text)
        self.assertIn("6　掃描資訊與範圍", text)
        self.assertIn("7　附錄", text)
        appendix = text.split("7　附錄")[-1]
        self.assertNotIn("AEO 問答檢測逐題結果", appendix)
        headings = [line for line in appendix.splitlines() if line.startswith("7.")]
        self.assertTrue(headings[-1].endswith("免責與報告產生方式"))

    def test_long_cve_lists_are_collapsed(self):
        cves = "、".join(f"CVE-2021-{4000 + i}" for i in range(20))
        self._finding("service-known-cve", "過時的 Apache", description=f"已知漏洞：{cves}。")
        problem = build_report_payload(self.scan_job)["findings"][0]["problem"]
        self.assertIn("CVE-2021-4007 等共 20 個 CVE", problem)
        self.assertNotIn("CVE-2021-4008", problem)

    def test_shared_component_vs_per_page_problem(self):
        for path in ("/", "/a", "/b"):
            self._finding("UX_TAP", "觸控目標過小", page=self._page(path),
                          evidence="過小的觸控目標：選單（30×17px）", category="ux")
            self._finding("geo-json-ld", "缺少結構化資料", page=Page.objects.get(
                scan_job=self.scan_job, url=f"https://example.com{path}"),
                evidence="json_ld_blocks=0", category="geo", severity="low")
        scopes = {f["title"]: f["scope"] for f in build_report_payload(self.scan_job)["findings"]}
        self.assertEqual(scopes["觸控目標過小"], "全站共用元件（3 個頁面出現同一處）")
        self.assertEqual(scopes["缺少結構化資料"], "影響 3 個頁面")

    def test_previous_scan_date_uses_full_date(self):
        from datetime import timedelta

        ScanJob.objects.create(
            user=self.user, original_url="https://example.com/",
            normalized_url="https://example.com/", origin="example.com",
            status=ScanJob.Status.COMPLETED, overall_score=60,
            completed_at=self.scan_job.completed_at - timedelta(days=3),
        )
        previous = build_report_payload(self.scan_job)["summary"]["previous"]
        self.assertRegex(previous["date"], r"^\d{4}-\d{2}-\d{2}$")
